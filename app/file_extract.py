"""Extracts plain text from tutor uploads so it can be fed to Claude."""

import io

import pdfplumber
from docx import Document


def extract_text(filename: str, content: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _extract_docx(content)
    if lower.endswith(".pdf"):
        return _extract_pdf(content)
    if lower.endswith(".txt"):
        return content.decode("utf-8")
    raise ValueError(f"Unsupported file type: {filename}")


def _extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" - ".join(cells))
    return "\n".join(lines)


def _extract_pdf(content: bytes) -> str:
    lines = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                for row in table:
                    cells = [c.strip() for c in row if c and c.strip()]
                    if cells:
                        lines.append(" - ".join(cells))
            if not page.extract_tables():
                text = page.extract_text()
                if text:
                    lines.append(text)
    return "\n".join(lines)
